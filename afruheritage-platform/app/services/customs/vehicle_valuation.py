from __future__ import annotations

from pathlib import Path
from typing import Any
import re
from difflib import SequenceMatcher

try:
    from openpyxl import load_workbook
except Exception:
    load_workbook = None


ROOT = Path(__file__).resolve().parents[3]
KRA_CRSP_FILE = ROOT / "data" / "customs" / "ke" / "kra_crsp_july_2025.xlsx"

GHANA_ICUMS_USED_VEHICLE_URL = (
    "https://external.unipassghana.com/cl/tm/tax/"
    "selectUsedVehicleTaxCalculate.do?MENU_ID=IIM01S03V02&decorator=popup"
)

GHANA_ICUMS_MODEL_URL = (
    "https://external.unipassghana.com/co/popup/"
    "selectCommonVehicleModelPopup.do"
)

GHANA_ICUMS_GENERAL_DUTY_URL = (
    "https://external.unipassghana.com/cl/tm/tax/"
    "selectGeneralTaxCalculate.do?MENU_ID=IIM01S03V01&decorator=popup"
)


def norm(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip().lower()




def sim(a: str, b: str) -> float:
    return SequenceMatcher(None, norm(a), norm(b)).ratio()


def token_score(query: str | None, target: str | None) -> float:
    if not query or not target:
        return 0.0

    q = norm(query)
    t = norm(target)

    if not q or not t:
        return 0.0

    if q == t:
        return 1.0

    if q in t or t in q:
        return 0.85

    return sim(q, t)


def detect_headers(rows: list[list[Any]]) -> tuple[int, dict[str, int]]:
    candidates = {
        "make": ["make", "manufacturer"],
        "model": ["model", "vehicle model"],
        "crsp": ["current retail selling price", "crsp"],
        "engine_cc": ["cc", "engine", "engine capacity"],
        "fuel": ["fuel"],
        "body": ["body", "type"],
    }

    for idx, row in enumerate(rows[:40]):
        headers = [norm(c) for c in row]
        found: dict[str, int] = {}

        for key, names in candidates.items():
            for i, h in enumerate(headers):
                if any(n in h for n in names):
                    found[key] = i
                    break

        if "make" in found and "model" in found and "crsp" in found:
            return idx, found

    return 0, {}


def load_kra_rows(limit: int = 5000) -> list[dict[str, Any]]:
    if load_workbook is None:
        raise RuntimeError("openpyxl is not installed")

    if not KRA_CRSP_FILE.exists():
        raise RuntimeError(f"KRA CRSP file not found: {KRA_CRSP_FILE}")

    wb = load_workbook(KRA_CRSP_FILE, data_only=True, read_only=True)
    ws = wb.active

    raw_rows = []
    for row in ws.iter_rows(values_only=True):
        raw_rows.append(list(row))
        if len(raw_rows) > limit + 50:
            break

    header_idx, headers = detect_headers(raw_rows)
    if not headers:
        raise RuntimeError("Could not detect KRA CRSP headers")

    items = []
    for row in raw_rows[header_idx + 1:]:
        if not any(row):
            continue

        def get(key):
            i = headers.get(key)
            return row[i] if i is not None and i < len(row) else None

        make = get("make")
        model = get("model")
        crsp = get("crsp")

        if not make or not model or not crsp:
            continue

        try:
            crsp_value = float(str(crsp).replace(",", ""))
        except Exception:
            continue

        items.append({
            "make": str(make).strip(),
            "model": str(model).strip(),
            "crsp_kes": crsp_value,
            "engine_cc": get("engine_cc"),
            "fuel_type": get("fuel"),
            "body_type": get("body"),
        })

        if len(items) >= limit:
            break

    return items


def search_kra_crsp(make: str | None, model: str | None, year: int | None = None) -> dict[str, Any]:
    rows = load_kra_rows()

    matches = []
    for row in rows:
        make_score = token_score(make, row.get("make"))
        model_score = token_score(model, row.get("model"))

        combined = (make_score * 0.40) + (model_score * 0.60)

        if make and not model:
            combined = make_score * 0.75

        if combined >= 0.45:
            candidate = dict(row)
            candidate["score"] = round(combined, 4)
            matches.append(candidate)

    matches.sort(key=lambda x: x["score"], reverse=True)

    if not matches:
        return {
            "matched": False,
            "source": "KRA CRSP July 2025",
            "message": "No KRA CRSP match found. Provide make and model, or connect a premium VIN decoder.",
            "matches": [],
        }

    best = matches[0]
    if best["score"] < 0.55:
        return {
            "matched": False,
            "source": "KRA CRSP July 2025",
            "message": "Low-confidence KRA CRSP match. Manual review required.",
            "best_candidate": best,
            "matches": matches[:10],
        }

    crsp = float(best["crsp_kes"])
    depreciation_rate = estimate_kenya_depreciation(year)
    customs_value = round(crsp * (1 - depreciation_rate), 2)

    return {
        "matched": True,
        "source": "KRA CRSP July 2025",
        "make": best["make"],
        "model": best["model"],
        "engine_cc": best.get("engine_cc"),
        "fuel_type": best.get("fuel_type"),
        "body_type": best.get("body_type"),
        "crsp_kes": crsp,
        "year": year,
        "estimated_depreciation_rate": depreciation_rate,
        "estimated_customs_value_kes": customs_value,
        "match_score": best["score"],
        "matches": matches[:10],
        "notes": [
            "Matched against official KRA CRSP July 2025 spreadsheet.",
            "Depreciation is estimated until the official valuation template is fully encoded.",
        ],
    }


def estimate_kenya_depreciation(year: int | None) -> float:
    if not year:
        return 0.0

    # Conservative approximation until full valuation template table is parsed.
    # Kenya import age limit is commonly up to 8 years; cap depreciation for safety.
    from datetime import datetime
    age = max(0, datetime.utcnow().year - int(year))

    if age <= 0:
        return 0.0
    if age == 1:
        return 0.10
    if age == 2:
        return 0.20
    if age == 3:
        return 0.30
    if age == 4:
        return 0.40
    if age == 5:
        return 0.50
    if age == 6:
        return 0.60
    if age == 7:
        return 0.65
    return 0.70


def ghana_icums_public_metadata() -> dict[str, Any]:
    return {
        "source": "Ghana ICUMS public pages",
        "used_vehicle_calculator": GHANA_ICUMS_USED_VEHICLE_URL,
        "vehicle_make_model_lookup": GHANA_ICUMS_MODEL_URL,
        "general_goods_duty_calculator": GHANA_ICUMS_GENERAL_DUTY_URL,
        "status": "adapter_ready",
        "notes": [
            "ICUMS public used vehicle calculator exposes VIN / make / model / year search.",
            "Result fields include HDV, currency, origin code, HS code, exchange rate, CIF NCY, and total tax.",
            "Next step is to build a compliant fetch/cache adapter around public form responses.",
        ],
    }


def estimate_ghana_from_vehicle_details(payload: dict[str, Any]) -> dict[str, Any]:
    year = payload.get("vehicle_year")
    make = payload.get("make")
    model = payload.get("model")
    engine_cc = payload.get("engine_cc")

    # Temporary reference estimate until ICUMS adapter/reference table is populated.
    base = 12000.0

    if year:
        age = max(0, 2026 - int(year))
        base = max(4000.0, base - (age * 650))

    if engine_cc:
        cc = int(engine_cc)
        if cc >= 3000:
            base *= 1.45
        elif cc >= 2000:
            base *= 1.25
        elif cc <= 1300:
            base *= 0.85

    freight = 1300.0
    insurance = round(base * 0.01, 2)
    hdv = round(base, 2)
    estimated_cif_usd = round(hdv + freight + insurance, 2)

    return {
        "matched": True,
        "source": "AfruHeritage Ghana interim reference valuation",
        "make": make,
        "model": model,
        "year": year,
        "engine_cc": engine_cc,
        "hdv_usd": hdv,
        "freight_usd": freight,
        "insurance_usd": insurance,
        "estimated_cif_usd": estimated_cif_usd,
        "notes": [
            "Interim Ghana valuation. Replace with ICUMS public adapter/cache when completed.",
            "Ghana ICUMS is known to generate HDV from VIN/origin/age, then apply insurance/freight to arrive at CIF.",
        ],
    }
