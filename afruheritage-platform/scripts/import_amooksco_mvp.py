import csv, hashlib, json, re, uuid
from pathlib import Path
import psycopg
from openpyxl import load_workbook

DB_URL = "postgresql://afruheritage:afruheritage@127.0.0.1:5433/afruheritage"
TENANT_NAME = "Amooksco Logistics"
TENANT_SLUG = "amooksco-logistics"
TENANT_EMAIL = "admin@amooksco.com"
BASE = Path("data/tenants/amooksco/imports")

def norm_phone(v):
    s = re.sub(r"\D+", "", str(v or ""))
    if not s:
        return ""
    if s.startswith("0") and len(s) == 10:
        s = "233" + s[1:]
    if len(s) == 9:
        s = "233" + s
    return s

def norm_name(v):
    return re.sub(r"\s+", " ", str(v or "").strip()).upper()

def tracking(seed):
    return "AFR-AMO-" + hashlib.sha1(seed.encode()).hexdigest()[:10].upper()

def ensure_tenant(conn):
    row = conn.execute("SELECT id FROM tenants WHERE slug=%s", (TENANT_SLUG,)).fetchone()
    if row:
        return row[0]

    tid = str(uuid.uuid4())
    conn.execute("""
        INSERT INTO tenants (
            id, company_name, slug, contact_email, plan_code,
            requested_domain, domain_type, launch_status,
            custom_domain, custom_domain_verified,
            created_at
        )
        VALUES (%s,%s,%s,%s,'free_trial',%s,'customer_domain','active',%s,false,now())
    """, (tid, TENANT_NAME, TENANT_SLUG, TENANT_EMAIL, "amooksco.com", "amooksco.com"))
    return tid

def import_whatsapp(conn, tenant_id, file_path):
    total = imported = duplicates = skipped = 0

    with open(file_path, newline="", encoding="utf-8-sig") as f:
        reader = csv.DictReader(f)
        print("CSV HEADERS:", reader.fieldnames)

        for r in reader:
            total += 1
            name = r.get("UserDisplayName") or r.get("Name") or r.get("display_name") or r.get("Display Name") or ""
            phone = r.get("UserPhone") or r.get("Phone") or r.get("phone") or r.get("Phone Number") or ""

            if not phone:
                for k, v in r.items():
                    if "phone" in str(k).lower() or "number" in str(k).lower():
                        phone = v
                        break

            nphone = norm_phone(phone)
            if not nphone:
                skipped += 1
                continue

            exists = conn.execute(
                "SELECT id FROM tenant_customers WHERE tenant_id=%s AND normalized_phone=%s",
                (tenant_id, nphone),
            ).fetchone()

            if exists:
                duplicates += 1
                continue

            conn.execute("""
                INSERT INTO tenant_customers (
                    tenant_id, display_name, normalized_name, phone, normalized_phone,
                    source, is_whatsapp_admin, is_whatsapp_super_admin, metadata
                )
                VALUES (%s,%s,%s,%s,%s,'whatsapp_csv',%s,%s,%s)
            """, (
                tenant_id,
                name,
                norm_name(name),
                phone,
                nphone,
                str(r.get("IsAdmin", "")).lower() in ("true", "1", "yes"),
                str(r.get("IsSuperAdmin", "")).lower() in ("true", "1", "yes"),
                json.dumps(r),
            ))
            imported += 1

    return {"total": total, "imported": imported, "duplicates": duplicates, "skipped": skipped}

def find_or_create_customer(conn, tenant_id, name, phone, source="manifest_shipping_mark"):
    nphone = norm_phone(phone)
    if nphone:
        row = conn.execute(
            "SELECT id FROM tenant_customers WHERE tenant_id=%s AND normalized_phone=%s",
            (tenant_id, nphone),
        ).fetchone()
        if row:
            return row[0], "phone_exact"

    clean_name = str(name or "").strip()
    nname = norm_name(clean_name)
    if nname:
        row = conn.execute(
            "SELECT id FROM tenant_customers WHERE tenant_id=%s AND normalized_name=%s LIMIT 1",
            (tenant_id, nname),
        ).fetchone()
        if row:
            return row[0], "name_exact"

        row = conn.execute("""
            INSERT INTO tenant_customers (
                tenant_id, display_name, normalized_name, phone, normalized_phone,
                source, metadata
            )
            VALUES (%s,%s,%s,%s,%s,%s,%s)
            RETURNING id
        """, (
            tenant_id,
            clean_name,
            nname,
            str(phone or ""),
            nphone or None,
            source,
            json.dumps({"created_from": source}),
        )).fetchone()

        return row[0], "created_from_manifest"

    return None, "unmatched"

def guess_columns(headers):
    mapping = {}
    for i, h in enumerate(headers):
        k = norm_name(h)
        if ("CUSTOMER" in k or "NAME" in k) and "customer_name" not in mapping:
            mapping["customer_name"] = i
        if ("PHONE" in k or "TEL" in k or "CONTACT" in k) and "phone" not in mapping:
            mapping["phone"] = i
        if "TRACK" in k and "external_tracking" not in mapping:
            mapping["external_tracking"] = i
        if "CBM" in k and "cbm" not in mapping:
            mapping["cbm"] = i
        if ("DESC" in k or "ITEM" in k or "GOODS" in k) and "description" not in mapping:
            mapping["description"] = i
        if "LOAD" in k and "DATE" in k and "loading_date" not in mapping:
            mapping["loading_date"] = i
        elif "DATE" in k and "receipt_date" not in mapping:
            mapping["receipt_date"] = i
    return mapping

def import_manifest(conn, tenant_id, file_path):
    wb = load_workbook(file_path, data_only=True)
    total_all = imported_all = matched_all = unmatched_all = errors_all = 0

    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue

        header_idx = None
        for idx, row in enumerate(rows[:10]):
            joined = " ".join(str(x or "") for x in row).upper()
            if "NAME" in joined or "CUSTOMER" in joined or "TRACK" in joined or "CBM" in joined:
                header_idx = idx
                break

        if header_idx is None:
            header_idx = 0

        headers = [str(x or "").strip() for x in rows[header_idx]]
        mapping = guess_columns(headers)
        print(f"XLSX {file_path.name} / sheet={ws.title} headers={headers}")
        print("MAPPING:", mapping)

        batch_id = conn.execute("""
            INSERT INTO cargo_import_batches (tenant_id, filename, source_type, status, metadata)
            VALUES (%s,%s,'excel_manifest','processing',%s)
            RETURNING id
        """, (tenant_id, f"{file_path.name}:{ws.title}", json.dumps({"headers": headers, "sheet": ws.title}))).fetchone()[0]

        total = imported = matched = unmatched = errors = 0

        for ridx, row in enumerate(rows[header_idx+1:], start=header_idx+2):
            if not row or all(v is None or str(v).strip() == "" for v in row):
                continue

            total += 1

            try:
                def get(col):
                    i = mapping.get(col)
                    return row[i] if i is not None and i < len(row) else None

                customer_name = get("customer_name") or row[0]
                phone = get("phone")
                external_tracking = get("external_tracking")
                description = get("description")
                cbm = get("cbm")
                receipt_date = get("receipt_date")
                loading_date = get("loading_date")

                customer_id, confidence = find_or_create_customer(conn, tenant_id, customer_name, phone)
                if customer_id:
                    matched += 1
                else:
                    unmatched += 1

                afru_tracking = tracking(f"{tenant_id}|{file_path.name}|{ws.title}|{ridx}|{customer_name}|{external_tracking}")

                conn.execute("""
                    INSERT INTO cargo_records (
                        tenant_id, customer_id, import_batch_id, afru_tracking_number,
                        external_tracking_number, customer_name, customer_phone, normalized_phone,
                        description, cbm, receipt_date, loading_date, status, match_confidence, raw_payload
                    )
                    VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,'received_at_warehouse',%s,%s)
                    ON CONFLICT (afru_tracking_number) DO NOTHING
                """, (
                    tenant_id, customer_id, batch_id, afru_tracking,
                    str(external_tracking or ""), str(customer_name or ""), str(phone or ""), norm_phone(phone),
                    str(description or ""), float(cbm) if isinstance(cbm, (int, float)) else None,
                    str(receipt_date or ""), str(loading_date or ""), confidence,
                    json.dumps({"row": ridx, "sheet": ws.title, "headers": headers, "values": [str(x) if x is not None else "" for x in row]}),
                ))

                cargo_id = conn.execute("SELECT id FROM cargo_records WHERE afru_tracking_number=%s", (afru_tracking,)).fetchone()[0]
                conn.execute("""
                    INSERT INTO cargo_tracking_events (tenant_id, cargo_record_id, status, message)
                    VALUES (%s,%s,'received_at_warehouse','Cargo record imported into Amooksco portal')
                """, (tenant_id, cargo_id))

                imported += 1

            except Exception as e:
                errors += 1
                print(f"ERROR row {ridx} in {file_path.name}/{ws.title}: {e}")

        conn.execute("""
            UPDATE cargo_import_batches
            SET total_rows=%s, imported_rows=%s, matched_rows=%s, unmatched_rows=%s,
                error_rows=%s, status='completed'
            WHERE id=%s
        """, (total, imported, matched, unmatched, errors, batch_id))

        total_all += total
        imported_all += imported
        matched_all += matched
        unmatched_all += unmatched
        errors_all += errors

    return {"file": file_path.name, "total": total_all, "imported": imported_all, "matched": matched_all, "unmatched": unmatched_all, "errors": errors_all}

def main():
    print("IMPORT DIR:", BASE.resolve())
    print("FILES:", [p.name for p in BASE.glob("*")])

    with psycopg.connect(DB_URL) as conn:
        tenant_id = ensure_tenant(conn)
        print("TENANT_ID", tenant_id)

        for f in BASE.rglob("*.csv"):
            print("CUSTOMERS", f.name, import_whatsapp(conn, tenant_id, f))

        for f in BASE.rglob("*.xlsx"):
            print("CARGO", import_manifest(conn, tenant_id, f))

        conn.commit()

        print("\nSUMMARY")
        for label, sql in [
            ("customers", "SELECT count(*) FROM tenant_customers WHERE tenant_id=%s"),
            ("cargo", "SELECT count(*) FROM cargo_records WHERE tenant_id=%s"),
            ("matched", "SELECT count(*) FROM cargo_records WHERE tenant_id=%s AND customer_id IS NOT NULL"),
            ("unmatched", "SELECT count(*) FROM cargo_records WHERE tenant_id=%s AND customer_id IS NULL"),
        ]:
            print(label, conn.execute(sql, (tenant_id,)).fetchone()[0])

if __name__ == "__main__":
    main()
